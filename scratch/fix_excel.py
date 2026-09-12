import openpyxl

file_path = r"E:\downloads chrome\sample-inventory.xlsx"
wb = openpyxl.load_workbook(file_path)
ws = wb.active

# 1. Fix "Purchase Rate " to "Purchase Rate"
for cell in ws[1]:
    if cell.value == "Purchase Rate ":
        cell.value = "Purchase Rate"

# Let's write the dropdown options manually in the helper columns if they are missing.
# Wait, if they used an ArrayFormula, I shouldn't mess with it unless it's missing values.
# But they specifically asked me to add the dropdowns for Category ("new or acid wash").
# They probably tried but couldn't get it fully working, or just wanted me to add it.
# Let's just overwrite the AO column with "New" and "Acid" starting from row 3 (since row 2 might be the array formula or a blank).
# Let's just see where Category options are stored.
# Earlier we saw:
# Col AO: ['Active Category Options', <ArrayFormula>, 'New', 0, None]
# This means AO3 is 'New'. Let's set AO4 to 'Acid'.
ws['AO4'] = 'Acid'

# Similarly for Type (Col AG is Type options):
# Let's make sure AG3 and AG4 are IN and OUT
ws['AG3'] = 'IN'
ws['AG4'] = 'OUT'

# For Hide Spec 1, 2, 3 (Col AQ):
ws['AQ3'] = 'Yes'
ws['AQ4'] = 'No'

# For Thread (Col AK):
ws['AK3'] = 'Full Thread'
ws['AK4'] = 'Half Thread'

wb.save("public/Book1- inventory.xlsx")
print("Saved modified user template to public/Book1- inventory.xlsx")
