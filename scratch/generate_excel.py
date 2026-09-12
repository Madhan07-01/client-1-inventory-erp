import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Stock Import Template"

headers = [
    "Product SKU / Size", "Warehouse", "Location / Rack", "Type", "Quantity",
    "Lot Number", "Brand", "Supplier", "Purchase Date", "Purchase Rate", 
    "Purchase Ref", "Grade", "Thread", "Finish", "Custom Spec 1", "Hide Spec 1",
    "Custom Spec 2", "Hide Spec 2", "Custom Spec 3", "Hide Spec 3",
    "Category", "Remarks"
]

ws.append(headers)

# Style headers
header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font

# Adjust column widths
for col in ws.columns:
    max_length = 0
    column = col[0].column_letter # Get the column name
    for cell in col:
        try: # Necessary to avoid error on empty cells
            if len(str(cell.value)) > max_length:
                max_length = len(cell.value)
        except:
            pass
    adjusted_width = (max_length + 2)
    ws.column_dimensions[column].width = adjusted_width

# Data Validation for Type (IN/OUT)
dv_type = DataValidation(type="list", formula1='"IN,OUT"', allow_blank=False)
ws.add_data_validation(dv_type)
dv_type.add(f"D2:D1000") # Column D is 'Type'

# Data Validation for Category (New/Acid)
dv_category = DataValidation(type="list", formula1='"New,Acid"', allow_blank=False)
ws.add_data_validation(dv_category)
# Category is index 20 (U) -> Column U
dv_category.add(f"U2:U1000")

# Data Validation for Hide Spec 1, 2, 3 (Yes/No)
dv_hide = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
ws.add_data_validation(dv_hide)
dv_hide.add(f"P2:P1000") # Hide Spec 1
dv_hide.add(f"R2:R1000") # Hide Spec 2
dv_hide.add(f"T2:T1000") # Hide Spec 3

# Add a couple of sample rows
sample1 = ["M20/70", "MADEENA MAIN", "BAY 1", "IN", 50, "LOT-001", "TVS", "ABC Fasteners", "2024-05-01", 15.50, "INV-8899", "8.8", "Half Thread", "Zinc", "", "No", "", "No", "", "No", "New", "Initial Stock"]
sample2 = ["M20/75", "MADEENA MAIN", "BAY 1", "IN", 100, "LOT-002", "Unbrako", "XYZ Corp", "2024-05-02", 16.20, "INV-9900", "10.9", "Full Thread", "HDG", "Special Coating", "Yes", "", "No", "", "No", "Acid", "Restock"]

ws.append(sample1)
ws.append(sample2)

wb.save("public/Book1- inventory.xlsx")
print("Successfully generated public/Book1- inventory.xlsx")
